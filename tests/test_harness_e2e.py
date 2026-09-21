"""
End-to-end tests for the Router-Free Unified Harness.
Tests workbook parsing, column mapping, building, validation, repair, and export.
"""

import os
import unittest
import openpyxl

from tools import (
    workbook_parser,
    column_mapper,
    excel_builder,
    control_plan_builder,
    validator,
    excel_exporter,
    ControlPlanRow
)
from harness import UnifiedHarness
from tests.generate_sample_pfmea import create_sample_pfmea


class TestUnifiedHarnessE2E(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.sample_path = "tests/data/sample_pfmea.xlsx"
        cls.output_path = "tests/data/test_output_Control_Plan.xlsx"
        create_sample_pfmea(cls.sample_path)

    def test_01_workbook_parser(self):
        rows = workbook_parser(self.sample_path)
        self.assertGreaterEqual(len(rows), 3)
        self.assertIn("Operation Number", rows[0])
        self.assertIn("Internal Work Center ID", rows[0])

    def test_02_column_mapper(self):
        rows = workbook_parser(self.sample_path)
        ctx = column_mapper(rows)

        self.assertEqual(len(ctx.items), 3)
        self.assertEqual(ctx.source_operations, ["10", "20", "30"])
        # Unmapped header check
        self.assertIn("Internal Work Center ID", ctx.unmapped_headers)
        # Check canonical mappings
        op10 = ctx.items[0]
        self.assertEqual(op10.operation_number, "10")
        self.assertEqual(op10.severity_rating, 7)
        self.assertEqual(op10.occurrence_rating, 4)

    def test_03_builder_and_validation(self):
        rows = workbook_parser(self.sample_path)
        ctx = column_mapper(rows)
        # Test both generic excel_builder and legacy alias
        built_rows = excel_builder(ctx, document_type="control_plan_from_pfmea")
        self.assertEqual(len(built_rows), 3)
        legacy_rows = control_plan_builder(ctx)
        self.assertEqual(len(legacy_rows), 3)

        # Rule V1: All 16 columns exist
        for r in built_rows:
            d = r.to_dict()
            self.assertEqual(len(d), 16)

        # Rule V3: Inventory numbers (tool_number, gage_number) must remain None on new import
        for r in built_rows:
            self.assertIsNone(r.tool_number)
            self.assertIsNone(r.gage_number)
            # Dynamic evidence-grounded extraction populates equipment or specs when mentioned
            if r.tool_name:
                self.assertIsInstance(r.tool_name, str)
                self.assertGreater(len(r.tool_name), 0)

        # Rule V5: Op 30 has no detective control -> reaction plan must be None
        op30_row = [r for r in built_rows if r.operation_number == "30"][0]
        self.assertIsNone(op30_row.reaction_plan)

        # Op 20 severity is 9 -> special_characteristic_class must be 'CC'
        op20_row = [r for r in built_rows if r.operation_number == "20"][0]
        self.assertEqual(op20_row.special_characteristic_class, "CC")

        # Validator check: Should be 0 violations
        violations = validator(built_rows, ctx)
        self.assertEqual(len(violations), 0, f"Expected 0 violations but got: {violations}")

    def test_04_repair_loop_on_violation(self):
        """Injects V3 and V5 violations and verifies the harness self-repairs."""
        rows = workbook_parser(self.sample_path)
        ctx = column_mapper(rows)
        built_rows = control_plan_builder(ctx)

        # Intentionally inject hallucination
        built_rows[0].tool_number = "TOOL-HALLUCINATED-99"
        built_rows[2].reaction_plan = "Hallucinated Reaction Plan without detective control"

        violations = validator(built_rows, ctx)
        self.assertGreater(len(violations), 0)

        # Harness repair
        harness = UnifiedHarness()
        repaired = harness._repair_rows(built_rows, violations, ctx, is_baseline=False)
        new_violations = validator(repaired, ctx)
        self.assertEqual(len(new_violations), 0, "Self-repair loop should resolve all injected violations.")

    def test_05_full_harness_run(self):
        harness = UnifiedHarness()
        log = harness.run(file_path=self.sample_path, output_path=self.output_path)

        self.assertEqual(log.status, "COMPLETE")
        self.assertEqual(log.output_rows_count, 3)
        self.assertEqual(len(log.violations), 0)
        self.assertTrue(os.path.exists(self.output_path))

        # Check the generated Excel file
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["Control Plan"]
        # Row 3 is header (16 columns)
        self.assertEqual(ws.max_column, 16)
        # Rows 4, 5, 6 are data rows
        self.assertGreaterEqual(ws.max_row, 6)
        wb.close()


if __name__ == "__main__":
    unittest.main()
