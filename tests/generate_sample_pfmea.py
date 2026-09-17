"""
Generates a realistic sample PFMEA .xlsx file for testing the harness.
Includes standard AIAG PFMEA columns, realistic manufacturing operations, and intentional unmapped columns.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def create_sample_pfmea(output_path: str = "tests/data/sample_pfmea.xlsx"):
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PFMEA"

    headers = [
        "Production Item Name",
        "Process Segment Name",
        "Operation Number",
        "Operation Name",
        "Product Characteristics",
        "Process Characteristics",
        "Potential Failure Mode",
        "Potential Effects of Failure",
        "Severity Rating: Failure Mode Effect",
        "Potential Causes of Failure",
        "Cause: Occurence Rating",
        "Preventive Controls",
        "Detective Controls",
        "Detective Controls: Rating",
        "Internal Work Center ID"  # Intentional unmapped header for audit test
    ]

    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    sample_rows = [
        # Op 10: CNC Rough Milling
        [
            "Steering Knuckle LH", "Machining", "10", "CNC Rough Milling",
            "Flange Thickness 15.0 +/- 0.2 mm", "Feed rate & Spindle RPM",
            "Machined face undersized", "Improper bolt seating in sub-assembly", 7,
            "Tool wear beyond threshold", 4,
            "Tool life management system with auto-offset",
            "CMM dimensional inspection every 2 hours", 3,
            "WC-MCH-01"
        ],
        # Op 20: Induction Hardening (High Severity -> CC)
        [
            "Steering Knuckle LH", "Heat Treatment", "20", "Induction Hardening",
            "Surface Hardness 58-62 HRC", "Quench water flow & coil temp",
            "Insufficient surface hardness", "Fatigue fracture of steering spindle", 9,
            "Quench nozzle clogged", 2,
            "Inline flow sensor with low-flow interlock",
            "Eddy current test 100% inline", 2,
            "WC-HT-04"
        ],
        # Op 30: Precision Bore Grinding (Preventive only, NO detective control -> V5 test)
        [
            "Steering Knuckle LH", "Grinding", "30", "Precision Bore Grinding",
            "Inner Bore Diameter 42.000 +0.015/-0.000 mm", "Dresser feed rate",
            "Bore taper exceeding tolerance", "Bearing premature seizure", 8,
            "Grinding wheel uneven wear", 3,
            "Automated CNC diamond wheel dressing every 25 parts",
            None,  # No detective control! Reaction plan must be null!
            None,
            "WC-GR-02"
        ]
    ]

    for row in sample_rows:
        ws.append(row)

    wb.save(output_path)
    wb.close()
    print(f"Sample PFMEA created at: {os.path.abspath(output_path)}")


if __name__ == "__main__":
    create_sample_pfmea()
