"""
Unit tests verifying layout decoupling between:
  1. Standard production single-sheet ('Control Plan') matching parent codebase.
  2. OEM APQP benchmark ('Header', 'Body') without bundled 'Control Plan'.
  3. Optional multi-tab bundle ('all') with both layouts.
"""

import os
import unittest
import openpyxl
from tools.models import ControlPlanRow
from tools.excel_exporter import excel_exporter


class TestLayoutModes(unittest.TestCase):

    def setUp(self):
        self.sample_rows = [
            ControlPlanRow(
                production_item_name="Machined Casing",
                process_segment_name="1",
                operation_number="10",
                operation_name="Rough Mill Flange",
                tool_number="T01",
                tool_name="CNC Mill",
                characteristic_id="1",
                process_characteristic="Spindle Speed 1200 RPM",
                product_characteristic="Flange Flatness",
                special_characteristic_class="CC",
                specification_tolerance="0.05 mm Max",
                control_method="CMM Inspection",
                gage_number="G-01",
                evaluation_measurement_technique="CMM",
                sample_size="3 pcs",
                sample_frequency="1/shift",
                reaction_plan="Quarantine and re-tool"
            ),
            ControlPlanRow(
                production_item_name="Machined Casing",
                process_segment_name="1",
                operation_number="20",
                operation_name="Finish Bore Hole",
                tool_number="T02",
                tool_name="Boring Bar",
                characteristic_id="2",
                process_characteristic="Feed Rate 0.1 mm/rev",
                product_characteristic="Bore Diameter",
                special_characteristic_class="SC",
                specification_tolerance="25.00 +/- 0.02 mm",
                control_method="Air Gage",
                gage_number="G-02",
                evaluation_measurement_technique="Air Gaging",
                sample_size="5 pcs",
                sample_frequency="Hourly",
                reaction_plan="Adjust tool offset"
            )
        ]

    def test_01_standard_export_produces_only_single_control_plan_sheet(self):
        out_path = "output/test_layout_standard.xlsx"
        res = excel_exporter(self.sample_rows, out_path, layout_mode="standard")
        self.assertTrue(os.path.exists(res))

        wb = openpyxl.load_workbook(res)
        self.assertEqual(wb.sheetnames, ["Control Plan"])
        ws = wb["Control Plan"]
        self.assertEqual(ws.max_column, 16)
        self.assertEqual(ws.max_row, 3)  # 1 header row + 2 data rows
        self.assertEqual(ws["A1"].value, "Production Item Name")
        self.assertEqual(ws["P1"].value, "Reaction Plan")
        # Validate parent codebase pale yellow header fill (#FFFFC5)
        self.assertIn("FFFFC5", str(ws["A1"].fill.start_color.rgb).upper())
        wb.close()

    def test_02_default_export_without_flag_is_single_control_plan_sheet(self):
        out_path = "output/test_layout_default.xlsx"
        res = excel_exporter(self.sample_rows, out_path)
        self.assertTrue(os.path.exists(res))

        wb = openpyxl.load_workbook(res)
        self.assertEqual(wb.sheetnames, ["Control Plan"])
        self.assertNotIn("Header", wb.sheetnames)
        self.assertNotIn("Body", wb.sheetnames)
        wb.close()

    def test_03_apqp_export_produces_only_header_and_body(self):
        out_path = "output/test_layout_apqp.xlsx"
        res = excel_exporter(self.sample_rows, out_path, layout_mode="apqp")
        self.assertTrue(os.path.exists(res))

        wb = openpyxl.load_workbook(res)
        self.assertEqual(wb.sheetnames, ["Header", "Body"])
        self.assertNotIn("Control Plan", wb.sheetnames)
        ws_b = wb["Body"]
        self.assertEqual(ws_b.max_column, 16)
        wb.close()

    def test_04_all_export_produces_bundle(self):
        out_path = "output/test_layout_bundle.xlsx"
        res = excel_exporter(self.sample_rows, out_path, layout_mode="all")
        self.assertTrue(os.path.exists(res))

        wb = openpyxl.load_workbook(res)
        self.assertIn("Control Plan", wb.sheetnames)
        self.assertIn("Header", wb.sheetnames)
        self.assertIn("Body", wb.sheetnames)
        wb.close()


if __name__ == "__main__":
    unittest.main()
